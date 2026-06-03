from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import INVENTORY_COLUMNS, PLAN_COLUMNS, THRESHOLDS
from .rules import AnalysisResult


HISTORY_DAYS = 30
INVENTORY_SHEET = "每次库存录入"
PLAN_SHEET = "每次进货记录"
PRODUCT_SHEET = "每次商品分析"
SUMMARY_SHEET = "每次风险汇总"
SEGMENT_SHEET = "30天商品分层建议"
LEGACY_SHEET_NAMES = {
    INVENTORY_SHEET: "每日库存录入",
    PLAN_SHEET: "每日进货记录",
    PRODUCT_SHEET: "每日商品分析",
    SUMMARY_SHEET: "每日风险汇总",
}

SUMMARY_COLUMNS = [
    "归档日期",
    "数据来源",
    "商品数",
    "异常商品数",
    "临期批次数",
    "快缺货",
    "库存偏高",
    "疑似滞销",
    "临期",
    "进货风险",
    "可能漏订",
]
PRODUCT_COLUMNS = [
    "归档日期",
    "数据来源",
    "商品键",
    "商品名",
    "条码",
    "品类",
    "上次进货总量",
    "库存剩余量",
    "单位",
    "本进货周期销量",
    "近30天销量",
    "预计可售进货周期数",
    "最早到期日",
    "已到期库存",
    "7天内临期库存",
    "8至30天临期库存",
    "货架位置",
    "仓库位置",
    "风险类型",
    "提示依据",
    "数据状态",
]
SEGMENT_COLUMNS = [
    "经营分层",
    "商品名",
    "品类",
    "归档次数",
    "30天累计进货量",
    "30天累计销量",
    "平均每周期销量",
    "最近库存剩余量",
    "最近预计可售进货周期数",
    "快缺货次数",
    "库存偏高次数",
    "疑似滞销次数",
    "临期风险次数",
    "建议动作",
    "判断依据",
]


def empty_history() -> dict[str, pd.DataFrame]:
    return {
        INVENTORY_SHEET: pd.DataFrame(columns=["归档日期", "数据来源", *INVENTORY_COLUMNS]),
        PLAN_SHEET: pd.DataFrame(columns=["归档日期", "数据来源", *PLAN_COLUMNS, "录入换算说明"]),
        PRODUCT_SHEET: pd.DataFrame(columns=PRODUCT_COLUMNS),
        SUMMARY_SHEET: pd.DataFrame(columns=SUMMARY_COLUMNS),
    }


def _date_text(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if pd.isna(value):
        return ""
    parsed = pd.to_datetime(value, errors="coerce")
    if pd.notna(parsed):
        return parsed.date().isoformat()
    return str(value).strip()


def _without_private_columns(frame: pd.DataFrame) -> pd.DataFrame:
    return frame.drop(
        columns=[column for column in frame.columns if str(column).startswith("_")],
        errors="ignore",
    ).copy()


def _add_archive_columns(frame: pd.DataFrame, analysis_date: date, source: str) -> pd.DataFrame:
    data = _without_private_columns(frame)
    data = data.drop(columns=["归档日期", "数据来源"], errors="ignore")
    data.insert(0, "数据来源", source)
    data.insert(0, "归档日期", analysis_date.isoformat())
    return data


def _summary_snapshot(
    result: AnalysisResult, analysis_date: date, source: str
) -> pd.DataFrame:
    row: dict[str, Any] = {
        "归档日期": analysis_date.isoformat(),
        "数据来源": source,
        "商品数": int(len(result.products)),
        "异常商品数": int(len(result.anomalies)),
        "临期批次数": int(len(result.expiry_batches)),
    }
    row.update({key: int(result.counts.get(key, 0)) for key in SUMMARY_COLUMNS[5:]})
    return pd.DataFrame([row], columns=SUMMARY_COLUMNS)


def _product_snapshot(
    result: AnalysisResult, analysis_date: date, source: str
) -> pd.DataFrame:
    products = _without_private_columns(result.products)
    risk_columns = ["商品键", "风险类型", "提示依据", "预计可售进货周期数"]
    if not result.anomalies.empty:
        risks = result.anomalies[risk_columns].copy()
    else:
        risks = pd.DataFrame(columns=risk_columns)
    products = products.drop(
        columns=["风险类型", "提示依据", "预计可售进货周期数"],
        errors="ignore",
    )
    products = products.merge(risks, on="商品键", how="left")
    products = _add_archive_columns(products, analysis_date, source)
    for column in PRODUCT_COLUMNS:
        if column not in products.columns:
            products[column] = ""
    return products[PRODUCT_COLUMNS]


def _replace_archive_date(
    existing: pd.DataFrame, new_rows: pd.DataFrame, analysis_date: date
) -> pd.DataFrame:
    date_text = analysis_date.isoformat()
    if existing.empty:
        return new_rows.reset_index(drop=True)
    data = existing.copy()
    if "归档日期" in data.columns:
        data = data[data["归档日期"].map(_date_text) != date_text]
    return pd.concat([data, new_rows], ignore_index=True)


def _normalise_dates(frame: pd.DataFrame) -> pd.DataFrame:
    data = frame.copy()
    if "归档日期" in data.columns:
        data["归档日期"] = data["归档日期"].map(_date_text)
    return data


def _latest_dates(history: dict[str, pd.DataFrame]) -> list[str]:
    dates: set[str] = set()
    for frame in history.values():
        if "归档日期" in frame.columns:
            dates.update(value for value in frame["归档日期"].map(_date_text) if value)
    return sorted(dates)[-HISTORY_DAYS:]


def trim_history(history: dict[str, pd.DataFrame]) -> dict[str, pd.DataFrame]:
    keep_dates = set(_latest_dates(history))
    if not keep_dates:
        return history
    trimmed: dict[str, pd.DataFrame] = {}
    for sheet, frame in history.items():
        data = _normalise_dates(frame)
        if "归档日期" in data.columns:
            data = data[data["归档日期"].isin(keep_dates)]
        trimmed[sheet] = data.reset_index(drop=True)
    return trimmed


def append_history_snapshot(
    history: dict[str, pd.DataFrame],
    result: AnalysisResult,
    inventory_frame: pd.DataFrame,
    plan_frame: pd.DataFrame | None,
    analysis_date: date,
    source: str,
) -> dict[str, pd.DataFrame]:
    archive = empty_history()
    for sheet, frame in history.items():
        if sheet in archive:
            archive[sheet] = _normalise_dates(frame)

    inventory = _add_archive_columns(inventory_frame, analysis_date, source)
    plan = (
        _add_archive_columns(plan_frame, analysis_date, source)
        if plan_frame is not None and not plan_frame.empty
        else pd.DataFrame(columns=archive[PLAN_SHEET].columns)
    )
    products = _product_snapshot(result, analysis_date, source)
    summary = _summary_snapshot(result, analysis_date, source)

    archive[INVENTORY_SHEET] = _replace_archive_date(
        archive[INVENTORY_SHEET], inventory, analysis_date
    )
    archive[PLAN_SHEET] = _replace_archive_date(archive[PLAN_SHEET], plan, analysis_date)
    archive[PRODUCT_SHEET] = _replace_archive_date(
        archive[PRODUCT_SHEET], products, analysis_date
    )
    archive[SUMMARY_SHEET] = _replace_archive_date(
        archive[SUMMARY_SHEET], summary, analysis_date
    )
    return trim_history(archive)


def read_history_workbook(content: bytes, filename: str) -> dict[str, pd.DataFrame]:
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("历史归档目前请上传系统导出的 .xlsx 文件。")
    try:
        sheets = pd.read_excel(BytesIO(content), sheet_name=None, dtype=object, engine="openpyxl")
    except Exception as exc:
        raise ValueError(f"历史归档读取失败：{exc}") from exc

    history = empty_history()
    for sheet in history:
        source_sheet = sheet if sheet in sheets else LEGACY_SHEET_NAMES.get(sheet, "")
        if source_sheet in sheets:
            history[sheet] = _normalise_dates(sheets[source_sheet].fillna(""))
    return trim_history(history)


def history_metrics(history: dict[str, pd.DataFrame]) -> dict[str, object]:
    dates = _latest_dates(history)
    products = history.get(PRODUCT_SHEET, pd.DataFrame())
    inventory = history.get(INVENTORY_SHEET, pd.DataFrame())
    latest_date = dates[-1] if dates else ""
    unique_products = 0
    if not products.empty and "商品键" in products.columns:
        unique_products = int(products["商品键"].replace("", pd.NA).dropna().nunique())
    elif not inventory.empty and "商品名" in inventory.columns:
        unique_products = int(inventory["商品名"].replace("", pd.NA).dropna().nunique())
    return {
        "归档天数": len(dates),
        "最新日期": latest_date or "暂无",
        "归档商品数": unique_products,
        "库存记录数": int(len(inventory)),
    }


def _number_or_zero(value: object) -> float:
    parsed = pd.to_numeric(pd.Series([value]), errors="coerce").iloc[0]
    return 0.0 if pd.isna(parsed) else float(parsed)


def _format_number(value: float | int | None) -> str:
    if value is None or pd.isna(value):
        return ""
    value = float(value)
    if value.is_integer():
        return str(int(value))
    return f"{value:.1f}"


def _risk_occurrences(values: pd.Series, label: str) -> int:
    return int(values.fillna("").astype(str).str.contains(label, regex=False).sum())


def history_product_segments(history: dict[str, pd.DataFrame]) -> pd.DataFrame:
    products = history.get(PRODUCT_SHEET, pd.DataFrame())
    if products.empty or "商品名" not in products.columns:
        return pd.DataFrame(columns=SEGMENT_COLUMNS)

    data = products.copy()
    for column in PRODUCT_COLUMNS:
        if column not in data.columns:
            data[column] = ""
    data["商品名"] = data["商品名"].fillna("").astype(str).str.strip()
    data = data[data["商品名"] != ""]
    if data.empty:
        return pd.DataFrame(columns=SEGMENT_COLUMNS)

    data["_归档日期"] = pd.to_datetime(data["归档日期"], errors="coerce")
    data["_商品分组键"] = data["商品键"].fillna("").astype(str).str.strip()
    data.loc[data["_商品分组键"] == "", "_商品分组键"] = data["商品名"]
    numeric_columns = [
        "上次进货总量",
        "库存剩余量",
        "本进货周期销量",
        "近30天销量",
        "预计可售进货周期数",
    ]
    for column in numeric_columns:
        data[column] = pd.to_numeric(data[column], errors="coerce")

    rows: list[dict[str, object]] = []
    for _, group in data.groupby("_商品分组键", sort=False):
        group = group.sort_values(["_归档日期", "归档日期"], kind="stable")
        latest = group.iloc[-1]
        archive_count = int(group["归档日期"].astype(str).replace("", pd.NA).dropna().nunique())
        if archive_count == 0:
            archive_count = int(len(group))

        risk_values = group["风险类型"]
        shortage_count = _risk_occurrences(risk_values, "快缺货")
        high_stock_count = _risk_occurrences(risk_values, "库存偏高")
        stagnant_count = _risk_occurrences(risk_values, "疑似滞销")
        expiry_count = _risk_occurrences(risk_values, "临期")

        purchase_total = float(group["上次进货总量"].fillna(0).sum())
        sales_total = float(group["本进货周期销量"].fillna(0).sum())
        valid_sales = group["本进货周期销量"].dropna()
        average_sales = float(valid_sales.mean()) if not valid_sales.empty else 0.0
        latest_stock = _number_or_zero(latest["库存剩余量"])
        latest_periods = latest["预计可售进货周期数"]
        if pd.isna(latest_periods) and average_sales > 0:
            latest_periods = latest_stock / average_sales

        numeric_periods = None if pd.isna(latest_periods) else float(latest_periods)
        if valid_sales.empty:
            segment = "数据不足"
            action = "继续按进货周期保存数据，至少累计 2 次后再做稳定判断。"
        elif archive_count < 2:
            segment = "观察中"
            action = "本次可作为参考；连续保存 2 次以上后再判断长期好卖或不好卖。"
        elif stagnant_count >= 2 or (sales_total <= 0 and latest_stock > 0):
            segment = "持续滞销"
            action = "建议减少或暂停进货，先做促销、换陈列或清理临期批次。"
        elif high_stock_count >= 2 or (
            numeric_periods is not None and numeric_periods > THRESHOLDS["库存偏高可售周期数"]
        ):
            segment = "库存偏高"
            action = "本轮谨慎进货；优先消化现有库存，避免越补越压货。"
        elif shortage_count >= 2 or (
            numeric_periods is not None and numeric_periods <= 1.5 and sales_total > 0
        ):
            segment = "稳定好卖"
            if numeric_periods is not None and numeric_periods <= THRESHOLDS["快缺货可售周期数"]:
                action = "建议优先补货，保持稳定供货；补货量要把现有库存一起扣进去。"
            else:
                action = "商品动销较好，但当前库存还能支撑，可按常规进货节奏补。"
        elif shortage_count == 1 or (
            numeric_periods is not None and numeric_periods <= 2 and sales_total > 0
        ):
            segment = "本周期需要补货"
            action = "建议列入本次进货复核，避免下个周期断货。"
        else:
            segment = "正常动销"
            action = "按常规进货周期维护，继续观察库存与销量变化。"

        periods_text = "无法消化" if average_sales <= 0 and latest_stock > 0 else _format_number(numeric_periods)
        evidence = (
            f"近 {archive_count} 次归档，累计销量 {_format_number(sales_total)}，"
            f"最近库存 {_format_number(latest_stock)}，预计可售 {periods_text or '待核对'} 个进货周期。"
        )
        rows.append(
            {
                "经营分层": segment,
                "商品名": latest["商品名"],
                "品类": latest.get("品类", ""),
                "归档次数": archive_count,
                "30天累计进货量": _format_number(purchase_total),
                "30天累计销量": _format_number(sales_total),
                "平均每周期销量": _format_number(average_sales),
                "最近库存剩余量": _format_number(latest_stock),
                "最近预计可售进货周期数": periods_text,
                "快缺货次数": shortage_count,
                "库存偏高次数": high_stock_count,
                "疑似滞销次数": stagnant_count,
                "临期风险次数": expiry_count,
                "建议动作": action,
                "判断依据": evidence,
            }
        )

    segments = pd.DataFrame(rows, columns=SEGMENT_COLUMNS)
    order = {
        "稳定好卖": 0,
        "本周期需要补货": 1,
        "库存偏高": 2,
        "持续滞销": 3,
        "正常动销": 4,
        "观察中": 5,
        "数据不足": 6,
    }
    if not segments.empty:
        segments["_排序"] = segments["经营分层"].map(order).fillna(99)
        segments["_销量排序"] = pd.to_numeric(segments["30天累计销量"], errors="coerce").fillna(0)
        segments = segments.sort_values(["_排序", "_销量排序"], ascending=[True, False], kind="stable")
        segments = segments.drop(columns=["_排序", "_销量排序"]).reset_index(drop=True)
    return segments


def build_history_workbook(
    history: dict[str, pd.DataFrame], exported_at: datetime
) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        notes = pd.DataFrame(
            [
                ("导出时间", exported_at.isoformat(timespec="seconds")),
                ("保留范围", f"最近 {HISTORY_DAYS} 个归档日期"),
                ("使用方式", "每个进货周期或每次分析后保存本次归档；下次上传该文件继续累计。"),
            ],
            columns=["项目", "内容"],
        )
        notes.to_excel(writer, index=False, sheet_name="使用说明")
        history_product_segments(history).to_excel(
            writer, index=False, sheet_name=SEGMENT_SHEET
        )
        for sheet, frame in empty_history().items():
            data = history.get(sheet, frame)
            if data.empty:
                data = pd.DataFrame(columns=frame.columns)
            data.to_excel(writer, index=False, sheet_name=sheet)

        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="147D68")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for column_cells in sheet.columns:
                letter = get_column_letter(column_cells[0].column)
                values = [str(cell.value or "") for cell in column_cells]
                sheet.column_dimensions[letter].width = max(
                    12, min(38, max(len(value) for value in values) * 1.5 + 3)
                )
    return output.getvalue()
