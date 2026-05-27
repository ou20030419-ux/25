from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import INVENTORY_COLUMNS, PLAN_COLUMNS


class InputReadError(ValueError):
    """上传文件无法按支持格式读取。"""


def read_table(content: bytes, filename: str) -> pd.DataFrame:
    lower_name = filename.lower()
    try:
        if lower_name.endswith(".xlsx"):
            return pd.read_excel(BytesIO(content), dtype=str, engine="openpyxl")
        if lower_name.endswith(".csv"):
            for encoding in ["utf-8-sig", "utf-8", "gb18030"]:
                try:
                    return pd.read_csv(BytesIO(content), dtype=str, encoding=encoding)
                except UnicodeDecodeError:
                    continue
            raise InputReadError("CSV 编码无法识别，请另存为 UTF-8 或 GB18030 后重试。")
    except InputReadError:
        raise
    except Exception as exc:
        raise InputReadError(f"文件读取失败：{exc}") from exc
    raise InputReadError("首版支持 .xlsx 与 .csv 文件，请更换文件格式后重试。")


def _sample_inventory() -> pd.DataFrame:
    return pd.DataFrame(
        [
            ["海盐薯片 60g", "2026-02-21", 3, "袋", 21, 75, 180, "", "A-01", "6900000000012", "膨化"],
            ["芝士饼干", "2026-02-13", 120, "盒", 14, 58, 300, "", "B-02", "6900000000029", "饼干"],
            ["芒果软糖", "2026-04-14", 26, "袋", 0, 8, 180, "", "C-03", "6900000000036", "糖果"],
            ["酸奶小蛋糕", "2026-05-21", 10, "个", 28, 96, 7, "", "D-01", "6900000000043", "糕点"],
            ["酸奶小蛋糕", "2026-05-19", 8, "个", 28, 96, 30, "", "D-01", "6900000000043", "糕点"],
            ["坚果能量棒", "2026-05-24", 2, "根", 14, 60, 180, "", "E-02", "6900000000050", "坚果"],
        ],
        columns=INVENTORY_COLUMNS,
    )


def _sample_plan() -> pd.DataFrame:
    return pd.DataFrame(
        [
            ["海盐薯片 60g", "6900000000012", 20, "袋", "周末陈列"],
            ["芝士饼干", "6900000000029", 40, "盒", ""],
            ["芒果软糖", "6900000000036", 12, "袋", ""],
            ["酸奶小蛋糕", "6900000000043", 20, "个", "核对临期批次"],
        ],
        columns=PLAN_COLUMNS,
    )


def template_frame(kind: str, sample: bool = False) -> pd.DataFrame:
    if kind == "inventory":
        return _sample_inventory() if sample else pd.DataFrame(columns=INVENTORY_COLUMNS)
    if kind == "plan":
        return _sample_plan() if sample else pd.DataFrame(columns=PLAN_COLUMNS)
    raise ValueError("未知模板类型")


def to_excel_bytes(frame: pd.DataFrame, sheet_name: str) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name=sheet_name)
        worksheet = writer.book[sheet_name]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="147D68")
            cell.alignment = Alignment(horizontal="center")
        for index, column in enumerate(frame.columns, start=1):
            width = max(12, min(26, len(str(column)) * 2 + 4))
            worksheet.column_dimensions[get_column_letter(index)].width = width
        if "条码" in frame.columns:
            barcode_column = frame.columns.get_loc("条码") + 1
            for row in range(2, max(worksheet.max_row + 1, 3)):
                worksheet.cell(row=row, column=barcode_column).number_format = "@"
    return output.getvalue()


def template_bytes(kind: str, sample: bool = False) -> bytes:
    sheet = "库存销售表" if kind == "inventory" else "计划进货单"
    return to_excel_bytes(template_frame(kind, sample), sheet)
