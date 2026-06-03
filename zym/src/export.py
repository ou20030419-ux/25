from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import RULESET_VERSION, SALES_BASE_LABEL, THRESHOLDS
from .rules import AnalysisResult


def _display_frame(frame: pd.DataFrame, columns: list[str] | None = None) -> pd.DataFrame:
    if frame.empty:
        return pd.DataFrame(columns=columns or list(frame.columns))
    shown = frame.copy()
    private = [column for column in shown.columns if str(column).startswith("_")]
    shown = shown.drop(columns=private, errors="ignore")
    return shown


def build_export_workbook(
    result: AnalysisResult,
    analysis_date: date,
    input_files: dict[str, str],
    exported_at: datetime,
) -> bytes:
    note_rows = [
        ("分析日期", analysis_date.isoformat()),
        ("导出时间", exported_at.isoformat(timespec="seconds")),
        ("规则版本", RULESET_VERSION),
        ("库存销售表文件", input_files.get("inventory", "")),
        ("计划进货单文件", input_files.get("plan", "未上传")),
        ("计划进货分析状态", "已执行" if result.plan_uploaded else "尚未上传计划进货单，未执行该分析"),
        ("销量计算基准", f"按一次进货周期录入销量；库存风险按{SALES_BASE_LABEL}估算"),
        ("库存剩余量算法", "优先使用手动填写的库存剩余量；未填写时按上次进货总量 - 本进货周期销量计算"),
        ("快缺货阈值", f"预计可售进货周期数 <= {THRESHOLDS['快缺货可售周期数']}"),
        ("库存偏高阈值", f"预计可售进货周期数 > {THRESHOLDS['库存偏高可售周期数']}"),
        ("可能进货过量阈值", f"进货后预计可售进货周期数 > {THRESHOLDS['进货过量可售周期数']}"),
        ("临期口径", "优先采用标注到期日期，否则按进货日期加保质期估算；仅统计有效且仍有库存的批次"),
        ("使用说明", "风险提示用于补货前复核，不代表强制性进货决定。"),
    ]
    notes = pd.DataFrame(note_rows, columns=["项目", "内容"])
    pending_note = pd.DataFrame(
        [{"说明": "尚未上传计划进货单，未执行该分析。"}]
    )
    plan_review = (
        _display_frame(result.plan_review)
        if result.plan_uploaded
        else pending_note
    )
    missed = (
        _display_frame(result.missed_orders)
        if result.plan_uploaded
        else pending_note
    )

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        sheets = {
            "分析说明": notes,
            "数据质量问题": _display_frame(result.quality_issues),
            "异常商品清单": _display_frame(result.anomalies),
            "临期批次明细": _display_frame(result.expiry_batches),
            "进货单自查结果": plan_review,
            "可能漏订清单": missed,
        }
        for name, frame in sheets.items():
            frame.to_excel(writer, index=False, sheet_name=name)
            sheet = writer.book[name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="147D68")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for column_cells in sheet.columns:
                letter = get_column_letter(column_cells[0].column)
                values = [str(cell.value or "") for cell in column_cells]
                sheet.column_dimensions[letter].width = max(
                    12, min(42, max(len(value) for value in values) * 1.7 + 3)
                )
    return output.getvalue()
